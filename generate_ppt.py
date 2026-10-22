import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    DARK_FOREST = RGBColor(16, 43, 40)      # #102B28
    EMERALD = RGBColor(27, 67, 50)          # #1B4332
    LIME_ACCENT = RGBColor(163, 230, 53)    # #A3E635
    GOLD_ACCENT = RGBColor(245, 158, 11)    # #F59E0B
    WARM_BG = RGBColor(251, 248, 243)       # #FBF8F3
    CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF
    TEXT_DARK = RGBColor(20, 35, 30)
    TEXT_MUTED = RGBColor(100, 116, 110)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(240, 243, 241)
    BORDER_COLOR = RGBColor(220, 226, 223)
    RED_ACCENT = RGBColor(220, 38, 38)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, section_number, title, subtitle=None):
        # Section Number & Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_sec = tf.paragraphs[0]
        p_sec.text = f"SECTION 0{section_number}" if isinstance(section_number, int) else section_number
        p_sec.font.name = "Segoe UI"
        p_sec.font.size = Pt(11)
        p_sec.font.bold = True
        p_sec.font.color.rgb = GOLD_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_FOREST

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Segoe UI"
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: Title Slide (Dark Forest Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_FOREST)

    # Accent decorative banner
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.8))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = LIME_ACCENT
    accent_bar.line.fill.background()

    # Title text box
    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "KANNADA DIGITAL INNOVATION HACKATHON 2026"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = LIME_ACCENT

    p1 = tf.add_paragraph()
    p1.text = "NammaMate (ನಮ್ಮ ಮೇಟ್)"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "A Little Kannada. A Lot More Connection."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(217, 249, 157)
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Track: Problem Statement 03 — Kannada in Everyday Life"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = WHITE

    p4 = tf.add_paragraph()
    p4.text = "Everyday Situational Assistant with AI Pronunciation Coach & Street Mode"
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(13)
    p4.font.color.rgb = RGBColor(180, 205, 195)

    # ==========================================
    # SLIDE 2: Problem Statement (Who, Why, Impact)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, WARM_BG)
    add_header(slide2, 1, "Problem Statement & Context", "Who is affected, why it matters, and the urban migration communication gap.")

    cards_data_s2 = [
        ("The Demographic Context", "Urban Influx in Karnataka",
         "Karnataka—particularly Bengaluru, Mysuru, Hubballi, and Mangaluru—welcomes millions of non-Kannada speaking students, tech engineers, and healthcare workers every year."),
        ("The Pain Point", "High-Stakes Language Anxiety",
         "When non-Kannada speakers step onto the street, communicate with auto drivers, negotiate at local vegetable markets, or visit pharmacies, intense fear of miscommunication leads to friction and awkwardness."),
        ("Why It Matters", "Cultural Harmony & Everyday Ease",
         "Language should build bridges, not barriers. Newcomers genuinely want to respect local customs, but conventional language tools fail to provide practical, conversational survival confidence.")
    ]

    for i, (tag, heading, body) in enumerate(cards_data_s2):
        left = Inches(0.8 + i * 3.95)
        top = Inches(2.0)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.8), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        ctb = slide2.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), Inches(3.2), Inches(4.1))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = tag.upper()
        cp0.font.name = "Segoe UI"
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = GOLD_ACCENT

        cp1 = ctf.add_paragraph()
        cp1.text = heading
        cp1.font.name = "Segoe UI"
        cp1.font.size = Pt(18)
        cp1.font.bold = True
        cp1.font.color.rgb = DARK_FOREST
        cp1.space_after = Pt(14)

        cp2 = ctf.add_paragraph()
        cp2.text = body
        cp2.font.name = "Segoe UI"
        cp2.font.size = Pt(13)
        cp2.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 3: Existing System & Limitations
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, WARM_BG)
    add_header(slide3, 2, "Existing Systems & Critical Gaps", "How everyday Kannada communication is handled today, and where existing tools fail.")

    gaps_data = [
        ("Google Translate / Generic LLMs",
         "• Generates overly formal, textbook Kannada that locals never use in daily speech.\n• Translates literal words without conversational context or cultural tone.\n• Unusable when you need an immediate response in a bustling auto stand."),
        ("Duolingo & Traditional Apps",
         "• Focuses on repetitive, academic grammar exercises ('The cat drinks milk').\n• Completely lacks high-frequency local survival scenarios like 'Meter haaki sir' or 'By-two coffee kodi'.\n• Requires weeks of commitment before a learner can speak a single sentence."),
        ("Printed Phrasebooks & Static Guides",
         "• Clunky and impossible to search quickly while on the move.\n• Zero native pronunciation feedback or syllable guidance.\n• Leads to silence, miscommunication, and friction in daily transactions.")
    ]

    for i, (title, points) in enumerate(gaps_data):
        top = Inches(1.9 + i * 1.7)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        # Red warning strip
        strip = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.1), Inches(1.5))
        strip.fill.solid()
        strip.fill.fore_color.rgb = RED_ACCENT
        strip.line.fill.background()

        tb = slide3.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = DARK_FOREST
        p0.space_after = Pt(4)

        for pt in points.split("\n"):
            p = tf.add_paragraph()
            p.text = pt
            p.font.name = "Segoe UI"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 4: How Your Solution Is Different
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, WARM_BG)
    add_header(slide4, 3, "How NammaMate Is Different", "Key innovations that elevate NammaMate into an indispensable everyday companion.")

    diff_features = [
        ("Street Mode / Driver Card", "High-contrast, extra-large Kannada cards designed to be tilted and shown directly to auto drivers and vendors when words escape you."),
        ("AI Pronunciation Coach", "Built-in speech recognition evaluates spoken Kannada with real-time phonetic match scores (e.g. 96%) and local praise ('Sakkath!')."),
        ("Interactive Roleplay AI", "Stepped conversation simulator (Bengaluru Auto ride, Darshini coffee ordering, KR Market bargaining) with scoring and etiquette tips."),
        ("Polite & Respectful Tone Engine", "Teaches crucial cultural etiquette—why using 'Banni'/'Kodi' instead of 'Ba'/'Kodu' wins local respect across Karnataka."),
        ("100% Offline-First Resilience", "All 50 phrases and audio tools are cached locally. Zero lag even in metro tunnels or low 5G connectivity areas."),
        ("One-Tap Emergency SOS", "Instant distress phrases with audio and verified 1-tap dialers for Karnataka 112 (Police), 108 (Ambulance), and 1091 (Women Safety).")
    ]

    for i, (title, desc) in enumerate(diff_features):
        col = i % 3
        row = i // 2
        left = Inches(0.8 + col * 3.95)
        top = Inches(2.0 + (i // 3) * 2.45)
        
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.8), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        accent = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.2), Inches(0.4), Inches(0.4))
        accent.fill.solid()
        accent.fill.fore_color.rgb = DARK_FOREST
        accent.line.fill.background()

        tb = slide4.shapes.add_textbox(left + Inches(0.7), top + Inches(0.18), Inches(2.9), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = DARK_FOREST
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 5: Approach & Methodology
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, WARM_BG)
    add_header(slide5, 4, "Approach & Key Design Decisions", "User-centric methodology designed for high-stress, real-world communication.")

    approach_steps = [
        ("Step 1: Situation First", "Divided by Real Needs", "Structured around the 5 highest-friction everyday pillars: Travel, Shopping, Food/Darshinis, Healthcare, and College."),
        ("Step 2: Bilingual Clarity", "Dual Script & Phonetics", "Every phrase displays clear Kannada script, Romanized English transliteration, natural meaning, and context tips."),
        ("Step 3: Practice & Verify", "AI Pronunciation Coach", "Users practice speaking with real-time phonetic string distance comparison, building vocal muscle memory before stepping out."),
        ("Step 4: Street Deployment", "Driver Card / Quick SOS", "If verbal speech fails under pressure, users flash the screen in Street Mode or trigger one-tap emergency dialers.")
    ]

    for i, (step, heading, desc) in enumerate(approach_steps):
        left = Inches(0.8 + i * 2.95)
        top = Inches(2.1)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(4.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        num_badge = slide5.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.25), top + Inches(0.3), Inches(0.6), Inches(0.6))
        num_badge.fill.solid()
        num_badge.fill.fore_color.rgb = DARK_FOREST
        num_badge.line.fill.background()

        tb_badge = slide5.shapes.add_textbox(left + Inches(0.25), top + Inches(0.3), Inches(0.6), Inches(0.6))
        tf_b = tb_badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = str(i + 1)
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(16)
        p_b.font.bold = True
        p_b.font.color.rgb = LIME_ACCENT
        p_b.alignment = PP_ALIGN.CENTER

        tb = slide5.shapes.add_textbox(left + Inches(0.25), top + Inches(1.1), Inches(2.3), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = step.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = GOLD_ACCENT

        p1 = tf.add_paragraph()
        p1.text = heading
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = DARK_FOREST
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 6: Tech Stack & System Architecture
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, WARM_BG)
    add_header(slide6, 5, "Tech Stack & Architecture Diagram", "Robust, modern, lightweight stack ensuring zero latency and offline capability.")

    # Left Column: Tech Stack Cards
    stack_items = [
        ("Frontend & UI Framework", "React 18 + TypeScript + Vite + Tailwind CSS (Responsive mobile-first)"),
        ("Audio & Speech Engine", "Web Speech Synthesis API ('kn-IN' dual-speed) & Web Speech Recognition API"),
        ("Scoring Algorithm", "Levenshtein String Distance & Phonetic Token Matcher for Pronunciation Feedback"),
        ("Cloud & Database", "Supabase PostgreSQL with Row Level Security & device-scoped synchronization"),
        ("Resilience & Cache", "Pre-cached Offline Dataset (offlineData.ts) + LocalStorage Fallback"),
    ]

    left_w = Inches(5.4)
    for i, (title, desc) in enumerate(stack_items):
        top = Inches(2.0 + i * 0.95)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, left_w, Inches(0.82))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tb = slide6.shapes.add_textbox(Inches(1.0), top + Inches(0.08), left_w - Inches(0.4), Inches(0.66))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = DARK_FOREST

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10)
        p1.font.color.rgb = TEXT_MUTED

    # Right Column: Architecture Diagram Visual Box
    diag_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2.0), Inches(6.0), Inches(4.7))
    diag_card.fill.solid()
    diag_card.fill.fore_color.rgb = DARK_FOREST
    diag_card.line.fill.background()

    diag_tb = slide6.shapes.add_textbox(Inches(6.8), Inches(2.2), Inches(5.4), Inches(4.3))
    dtf = diag_tb.text_frame
    dtf.word_wrap = True

    dp0 = dtf.paragraphs[0]
    dp0.text = "SYSTEM ARCHITECTURE & DATA FLOW"
    dp0.font.name = "Segoe UI"
    dp0.font.size = Pt(13)
    dp0.font.bold = True
    dp0.font.color.rgb = LIME_ACCENT
    dp0.space_after = Pt(12)

    diag_text = (
        "┌──────────────────────────────────────────────────┐\n"
        "│              CLIENT USER INTERFACE               │\n"
        "│  [Situations]  [Roleplay AI]  [Street Mode Card] │\n"
        "│  [Smart Search] [Favorites]   [SOS Emergency]    │\n"
        "└────────────────────────┬─────────────────────────┘\n"
        "                         │\n"
        "         ┌───────────────┴───────────────┐\n"
        "         ▼                               ▼\n"
        "┌───────────────────┐           ┌──────────────────┐\n"
        "│  WEB SPEECH API   │           │ OFFLINE CACHE    │\n"
        "│ - SpeechSynthesis │           │ (Zero-Lag 50     │\n"
        "│ - kn-IN SpeechRec │           │  Survival Words) │\n"
        "└────────┬──────────┘           └────────┬─────────┘\n"
        "         │                               │\n"
        "         └───────────────┬───────────────┘\n"
        "                         ▼\n"
        "┌──────────────────────────────────────────────────┐\n"
        "│             SUPABASE CLOUD DATABASE             │\n"
        "│  PostgreSQL • Row Level Security • Device Sync   │\n"
        "└──────────────────────────────────────────────────┘"
    )

    dp1 = dtf.add_paragraph()
    dp1.text = diag_text
    dp1.font.name = "Consolas"
    dp1.font.size = Pt(9.5)
    dp1.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 7: Conclusion, Impact & Future Scope
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, WARM_BG)
    add_header(slide7, 6, "Impact, Conclusion & Future Scope", "Bridging cultural gaps and turning hesitation into connection.")

    conc_cards = [
        ("Real-World Impact", "Promoting Cultural Harmony",
         "• Eliminates fear and tension during daily transit and market visits.\n• Shows non-Kannada speakers that learning Kannada is warm, approachable, and fun.\n• Builds immediate goodwill with auto drivers, vendors, and locals through respectful etiquette."),
        ("Future Roadmap", "Scaling Across Karnataka",
         "• Regional Dialect Explorer: Expanding into North Karnataka, Kundapura/Mangaluru, and Mysuru colloquial idioms.\n• Verified Native Audio: Crowdsourcing natural local voice clips from students across Karnataka.\n• WhatsApp/Telegram Quick Bot: Instant survival phrase lookup directly inside chat apps.")
    ]

    for i, (title, sub, body) in enumerate(conc_cards):
        left = Inches(0.8 + i * 5.95)
        top = Inches(2.0)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(3.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tb = slide7.shapes.add_textbox(left + Inches(0.35), top + Inches(0.3), Inches(5.05), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = GOLD_ACCENT

        p1 = tf.add_paragraph()
        p1.text = sub
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = DARK_FOREST
        p1.space_after = Pt(10)

        for line in body.split("\n"):
            p = tf.add_paragraph()
            p.text = line
            p.font.name = "Segoe UI"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(4)

    # Bottom Callout Banner
    banner = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = DARK_FOREST
    banner.line.fill.background()

    btb = slide7.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.9))
    btf = btb.text_frame
    btf.word_wrap = True

    bp0 = btf.paragraphs[0]
    bp0.text = "Kannada Kaliyona, Karnataka Preetisana! (ಕನ್ನಡ ಕಲಿಯೋಣ, ಕರ್ನಾಟಕ ಪ್ರೀತಿಸೋಣ!)"
    bp0.font.name = "Segoe UI"
    bp0.font.size = Pt(16)
    bp0.font.bold = True
    bp0.font.color.rgb = LIME_ACCENT

    bp1 = btf.add_paragraph()
    bp1.text = "Live Demo & Source Code: Fully deployed on web • Open-source on GitHub"
    bp1.font.name = "Segoe UI"
    bp1.font.size = Pt(12)
    bp1.font.color.rgb = WHITE

    output_path = os.path.join(os.path.dirname(__file__), "NammaMate_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
